import re
from datetime import date, timedelta
from io import BytesIO
from openpyxl import load_workbook


def parse_excel(content):
    workbook = load_workbook(BytesIO(content), data_only=True)
    weeks = []
    for sheet in workbook:
        match = re.fullmatch(r'S[IP]-(\d{2})', sheet.title.strip())
        if not match:
            continue
        number = int(match[1])
        title = str(sheet['A1'].value or '').strip()
        years = re.findall(r'20\d{2}', title)
        if not years:
            raise ValueError(f'Année introuvable : {sheet.title}')
        # The final year can be next year for ISO week 53 crossing New Year.
        year = int(years[0])
        if number >= 52 and 'JANVIER' in title.upper():
            year = int(years[-1]) - 1
        if number == 1 and len(years) > 1:
            year = int(years[-1])
        monday = date.fromisocalendar(year, number, 1)
        days, current = [], None
        for row in sheet.iter_rows():
            values = [str(c.value).strip() if c.value is not None else '' for c in row]
            values += [''] * max(0, 12 - len(values))
            label = next((v.upper() for v in values[7:9] if v.upper() in ['LUNDI', 'MARDI', 'MERCREDI', 'JEUDI', 'VENDREDI', 'SAMEDI', 'DIMANCHE']), '')
            if label in ['LUNDI', 'MARDI', 'MERCREDI', 'JEUDI', 'VENDREDI', 'SAMEDI', 'DIMANCHE']:
                offset = ['LUNDI', 'MARDI', 'MERCREDI', 'JEUDI', 'VENDREDI', 'SAMEDI', 'DIMANCHE'].index(label)
                current = {'name': label.capitalize(), 'date': (monday + timedelta(days=offset)).isoformat(), 'routes': []}
                days.append(current)
            if current and any(values[:7]) and not (values[0].upper().startswith('SEMAINE')):
                current['routes'].append(dict(zip(['number', 'route', 'type', 'vehicle', 'driver', 'crew1', 'crew2'], values[:7])))
        if not days or not any(day['routes'] for day in days):
            raise ValueError(f'Aucune tournée : {sheet.title}')
        weeks.append({'id': f'{year}-W{number:02}', 'number': number, 'year': year, 'title': title, 'start': monday.isoformat(), 'days': days})
    workbook.close()
    if not weeks:
        raise ValueError('Aucune feuille hebdomadaire SI-xx ou SP-xx reconnue.')
    if len({w['id'] for w in weeks}) != len(weeks):
        raise ValueError('Semaines en double.')
    return sorted(weeks, key=lambda week: week['start'])
